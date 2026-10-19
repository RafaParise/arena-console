#!/usr/bin/env python3
"""GERADOR DE ATIVIDADES DIDATICAS ARENA - 6 aulas praticas em .docx"""
import os, sys
try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("ERRO: python-docx nao instalado. Use Python 3.14.")
    sys.exit(1)

PASTA = os.getcwd()
COR_AZUL = RGBColor(0x00, 0x66, 0xCC)
COR_VERDE = RGBColor(0x2E, 0x7D, 0x38)
COR_LAR = RGBColor(0xFF, 0x66, 0x00)

# === CODIGO DOS SNIPPETS (triple-quoted, no escaping) ===
TLS1 = """if partes.scheme != "https":
    raise ValueError("Use HTTPS")"""

TLS2 = """if digest.hexdigest() != sha256:
    raise ValueError("Arquivo corrompido")"""

TLS3 = """PORTA = 8777  # servidor HTTP, sem TLS
# Na rede local, apelidos e recordes trafegam abertos"""

SCHEMA = """CREATE TABLE placar (
    uuid TEXT PRIMARY KEY,
    apelido TEXT NOT NULL,
    jogo_id TEXT NOT NULL,
    pontos INTEGER NOT NULL,
    registrado_em TEXT NOT NULL
);
CREATE TABLE sessao (
    uuid TEXT PRIMARY KEY,
    apelido TEXT NOT NULL,
    jogo_id TEXT NOT NULL,
    inicio TEXT NOT NULL,
    segundos INTEGER NOT NULL
);"""

RANKING = """SELECT * FROM (
    SELECT jogo_titulo AS jogo,
           apelido AS apelido,
           MAX(pontos) AS pontos,
           RANK() OVER (
               PARTITION BY jogo_id
               ORDER BY MAX(pontos) DESC
           ) AS posicao
    FROM placar
    GROUP BY jogo_id, apelido
)
WHERE posicao <= 5
ORDER BY jogo, posicao;"""

HTML_SNIPPET = """<main class="corpo">
  <nav class="trilha">...</nav>
  <section class="conteudo">
    <ul id="listaJogos"></ul>
  </section>
  <aside class="detalhes-jogo">
    <button id="btnJogarSelecionado">Jogar agora</button>
  </aside>
</main>"""

ESTADO_JS = """const estado = {
  jogos: [],
  config: null,
  jogoSelecionado: null,
};  // fonte de verdade so
const $ = (seletor) => document.querySelector(seletor);
// $() e como chamar alguem pelo nome"""

PEDIR_JS = """async function pedir(rota, metodo = 'GET', corpo = null) {
  const opcoes = { method: metodo, headers: { 'Content-Type': 'application/json' } };
  if (corpo) opcoes.body = JSON.stringify(corpo);
  let resposta;
  try {
    resposta = await fetch(rota, opcoes);
  } catch (falha) {
    perderContato();
    throw new Error('O ARENA parou de responder...');
  }
  const dados = await resposta.json();
  return dados;
}"""

EVENTOS_JS = """$('#btnJogarSelecionado').addEventListener('click', function() {
    var jogo = estado.jogoSelecionado;
    if (jogo) jogar(jogo);
})"""

DOM_COMANDOS = """// TESTE 1: Ler um elemento
$('#versao').textContent
$('#detalhesTitulo').textContent

// TESTE 2: Mudar um texto
$('#detalhesTitulo').textContent = 'Turma do ARENA!'
$('#detalhesTitulo').style.color = 'blue'

// TESTE 3: Esconder e mostrar
$('#detalhesCapa').hidden = true
$('#detalhesCapa').hidden = false

// TESTE 4: Criar novo elemento
var aviso = document.createElement('p')
aviso.textContent = 'ATIVIDADE CONCLUIDA!'
aviso.style.background = 'yellow'
$('#detalhesCorpo').prepend(aviso)

// TESTE 5: Registrar evento
var b = document.querySelector('#btnJogarSelecionado')
b.addEventListener('click', function() {
    alert('Este botao abre um jogo pelo emulador!')
})"""

API_COMANDOS = """// Pedido 1: quem esta online?
fetch('/api/colegas').then(r => r.json()).then(d => console.log(d))

// Pedido 2: ranking de pontos
fetch('/api/ranking').then(r => r.json()).then(d => console.log(d))

// Pedido 3: conquistas
fetch('/api/conquistas').then(r => r.json()).then(d => console.log(d))

// Pedido 4: versao do servidor
fetch('/api/estado').then(r => r.json()).then(d => console.log(d.versao))"""

# === FUNCOES AUXILIARES ===
def novodoc():
    doc = Document()
    for n, sz in [("Title", 20), ("Heading 1", 14), ("Normal", 11)]:
        s = doc.styles[n]
        s.font.name = "Calibri"
        s.font.size = Pt(sz)
        s.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    doc.styles["Title"].font.bold = True
    doc.styles["Title"].font.color.rgb = COR_AZUL
    doc.styles["Heading 1"].font.color.rgb = COR_AZUL
    return doc

def titulo(d, t):
    p = d.add_paragraph(t, style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_after = Pt(12)

def subtitulo(d, t):
    p = d.add_paragraph(t, style="Heading 1")
    p.space_before = Pt(12)
    p.space_after = Pt(6)

def texto(d, t):
    p = d.add_paragraph(t)
    p.space_after = Pt(6)

def etiqueta(d, t, cor=None):
    p = d.add_paragraph()
    r = p.add_run(t)
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.bold = True
    if cor: r.font.color.rgb = cor
    p.space_after = Pt(3)

def listinha(d, items):
    for item in items:
        p = d.add_paragraph(item, style="List Bullet")
        p.style.font.name = "Calibri"
        p.style.font.size = Pt(11)

def comando(d, codigo):
    p = d.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.right_indent = Inches(0.3)
    for line in codigo.split("\n"):
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(9)
        if line != codigo.split("\n")[-1]:
            r.add_break()
    p.space_after = Pt(8)

def caixa(d, titulo, texto, emoji="[i]"):
    p = d.add_paragraph()
    p.space_before = Pt(4)
    p.space_after = Pt(4)
    r = p.add_run(emoji + " " + titulo)
    r.font.name = "Calibri"
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = COR_LAR
    p2 = d.add_paragraph()
    p2.paragraph_format.left_indent = Inches(0.3)
    p2.paragraph_format.right_indent = Inches(0.3)
    for line in texto.split("\n"):
        rr = p2.add_run(line)
        rr.font.name = "Calibri"
        rr.font.size = Pt(10)
        if line != texto.split("\n")[-1]:
            rr.add_break()
    p2.space_after = Pt(4)

def tabreg(d, headers, rows):
    t = d.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.cell(0, i)
        c.text = h
        s = c._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:fill"), "0066CC")
        c._tc.append(sh)
        for r in c.paragraphs[0].runs:
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row in rows:
        ln = t.add_row()
        for i, val in enumerate(row):
            cl = ln.cells[i]
            cl.text = str(val)
            for r in cl.paragraphs[0].runs:
                r.font.name = "Calibri"
                r.font.size = Pt(10)
    d.add_paragraph()

def checklist(d, items):
    for item in items:
        p = d.add_paragraph("[ ] " + item, style="List Bullet")
        p.style.font.name = "Calibri"
        p.style.font.size = Pt(11)

def limpar():
    for f in os.listdir(PASTA):
        if f.startswith("ATIVIDADE-ARENA-AULA") and f.endswith(".docx"):
            os.remove(os.path.join(PASTA, f))

def gerar(nome, titulo_principal, subtitulo_texto, secoes):
    d = novodoc()
    titulo(d, titulo_principal)
    texto(d, subtitulo_texto)
    for bloco in secoes:
        tp = bloco[0]
        if tp == "obj":
            subtitulo(d, "OBJETIVO")
            texto(d, bloco[1])
        elif tp == "mat":
            subtitulo(d, "MATERIAL")
            listinha(d, bloco[1])
        elif tp == "passos":
            subtitulo(d, "PASSO A PASSO")
            for i, (prof, alun) in enumerate(bloco[1], 1):
                etiqueta(d, "PASSO " + str(i))
                if prof:
                    etiqueta(d, "PROFESSOR:", COR_AZUL)
                    listinha(d, prof)
                if alun:
                    etiqueta(d, "ALUNO:", COR_VERDE)
                    listinha(d, alun)
        elif tp == "cmd":
            subtitulo(d, "COMANDOS")
            comando(d, bloco[1])
        elif tp == "ana":
            caixa(d, "Analogia", bloco[1])
        elif tp == "cod":
            subtitulo(d, "CODIGO DO ARENA")
            texto(d, bloco[1])
            comando(d, bloco[2])
        elif tp == "tab":
            subtitulo(d, bloco[1])
            hdrs = bloco[2][0]
            if isinstance(bloco[2][1], list) and bloco[2][1] and isinstance(bloco[2][1][0], list):
                rws = bloco[2][1]
            else:
                rws = list(bloco[2][1:])
            tabreg(d, hdrs, rws)
        elif tp == "info":
            caixa(d, "Dica", bloco[1])
        elif tp == "perg":
            subtitulo(d, "PERGUNTAS")
            listinha(d, bloco[1])
            subtitulo(d, "RESPOSTA")
            for r in bloco[2]:
                texto(d, r)
        elif tp == "res":
            subtitulo(d, "RESULTADO ESPERADO")
            checklist(d, bloco[1])
        elif tp == "saida":
            subtitulo(d, "BILHETE")
            texto(d, bloco[1])
        elif tp == "txt":
            texto(d, bloco[1])
    texto(d, "Desenvolvido com o ARENA Console - sem instalar, sem "
               "internet, sem administrador.")
    c = os.path.join(PASTA, nome)
    d.save(c)
    print("  OK " + nome + " (" + str(os.path.getsize(c)) + " bytes)")


# ============================================================
# AULAS
# ============================================================

def aula1():
    gerar(
        "ATIVIDADE-ARENA-AULA1-REDES.docx",
        "ATIVIDADE: REDES - TLS 1.3",
        "Como o ARENA protege os downloads de jogos",
        [
            ("obj",
             "Entender as 3 garantias do TLS (confidencialidade, integridade, "
             "autenticidade) e identificar o papel do cliente e do provedor "
             "numa conexao HTTPS."),
            ("mat",
             ["Notebook com ARENA.bat aberto",
              "Navegador (Chrome ou Edge)",
              "Quadro branco ou lousa"]),
            ("passos", [
                (["Escreva na lousa: TLS protege 3 coisas:",
                  "1) Confidencialidade (ninguem le)",
                  "2) Integridade (ninguem altera)",
                  "3) Autenticidade (o site e real)"],
                 None),
                (["Mostre os 3 trechos de codigo reais do ARENA",
                  "Divida a turma em grupos"],
                 ["Identifique provedor vs cliente em cada trecho",
                  "Anote a garantia de seguranca que cada trecho protege"]),
                (["Abra http://127.0.0.1:8777/api/ranking no navegador",
                  "Pergunte: esses dados sao pessoais?"],
                 ["Leia o JSON", "Anote dados que identificam uma pessoa"]),
            ]),
            ("ana",
             "O TLS 1.3 e como um envelope lacrado. Voce baixa um jogo do "
             "archive.org. O TLS garante que ninguem leu (confidencialidade), "
             "alterou (integridade) ou falsificou (autenticidade) o pacote. "
             "No TLS 1.2 eram duas idas e voltas; no 1.3 basta uma — mais "
             "rapido e seguro."),
            ("cod",
             "Trecho 1 - acervo.py (linha 79): forca HTTPS",
             TLS1 + "\n\n"
             "Trecho 2 - acervo.py (linha 226): verifica SHA-256",
             TLS2 + "\n\n"
             "Trecho 3 - arena.py (linha 133): HTTP sem TLS",
             TLS3),
            ("info",
             "Provedor: certificado valido, TLS 1.3 ligado.\n"
             "Cliente: usa https, nao ignora alertas, conferir SHA-256.\n"
             "No ARENA, aluno e cliente ao baixar jogos (HTTPS) e provedor "
             "ao hospedar partidas (HTTP local). LGPD protege apelidos no ranking."),
            ("res",
             ["Entendeu as 3 garantias do TLS",
              "Identificou provedor vs cliente nos trechos",
              "Reconheceu dados pessoais no ranking (LGPD)"]),
            ("saida",
             "Cite uma responsabilidade do provedor e uma do cliente."),
        ]
    )

def aula2():
    gerar(
        "ATIVIDADE-ARENA-AULA2-REDES-PRATICA.docx",
        "ATIVIDADE: REDES - Testando TLS no Laboratorio",
        "Testando o handshake TLS 1.3 com OpenSSL",
        [
            ("obj",
             "Executar testes de handshake TLS, observar a recusa de versoes "
             "antigas e analisar o risco da exposicao de dados na rede local."),
            ("mat",
             ["Notebook com ARENA.bat aberto",
              "Wi-Fi da escola ou hotspot do celular",
              "Arquivo TESTAR_TLS.bat"]),
            ("passos", [
                (["Anote o IP do notebook na tela (ex: 192.168.0.10:8777)",
                  "Explique que TESTAR_TLS.bat usa OpenSSL embutido"],
                 ["Abra TESTAR_TLS.bat com dois cliques",
                  "Janela preta mostra versao OpenSSL"]),
                (None,
                 ["Leia os resultados:",
                  "Teste 1: 3 sites com TLS 1.3",
                  "Teste 2: TLS 1.2 (compatibilidade)",
                  "Teste 3: TLS 1.1 (deve ser recusado)"]),
                (["Abra http://127.0.0.1:8777/api/ranking",
                  "Mostre que dados do ranking sao pessoais"],
                 ["Copie os dados da tabela"]),
            ]),
            ("ana",
             "Teste 1: sites confiaveis usam TLS 1.3.\n"
             "Teste 2: compatibilidade mas menos segura.\n"
             "Teste 3: cliente recusa TLS 1.1 — como carteiro sem documento."),
            ("tab",
             "Ficha de registro - TLS",
             (["Site", "Versao TLS", "Cifra", "Emissor", "Valido ate"],
              [["archive.org", "", "", "", ""],
               ["raw.githubusercontent.com", "", "", "", ""],
               ["github.com (teste 1)", "", "", "", ""],
               ["github.com (teste 2)", "", "", "", ""],
               ["github.com (teste 3)", "", "", "", ""]])),
            ("info",
             "getaddrinfo = sem internet = ligue hotspot\n"
             "CERTIFICATE_VERIFY_FAILED = protecao funcionando\n"
             "Emissor Zscaler = rede inspeciona HTTPS"),
            ("res",
             ["OpenSSL 3.0.21+ rodando",
              "TLS 1.3 passa, TLS 1.2 aceito, TLS 1.1 recusado",
              "Identificou dados pessoais no ranking local"]),
            ("saida",
             "Por que o teste 3 foi recusado? Provedor ou cliente decidiu?"),
        ]
    )

def aula3():
    gerar(
        "ATIVIDADE-ARENA-AULA3-BANCO.docx",
        "ATIVIDADE: BANCO DE DADOS - SQLite + SQL",
        "Explorando o banco de dados do ARENA",
        [
            ("obj",
             "Utilizar comandos SQL de agregacao (GROUP BY, RANK, PARTITION BY) "
             "sobre uma base real e interpretar resultados na tela."),
            ("mat",
             ["Python do ARENA (servidor/python/python.exe)",
              "Arquivo dados/arena.db",
              "Arquivo TESTAR_BANCO.bat"]),
            ("passos", [
                (["Explique: ARENA guarda recordes num arquivo .db (SQLite)",
                  "SQLite e um banco em arquivo, sem instalar nada"],
                 ["Abra TESTAR_BANCO.bat",
                  "Leia: 3 tabelas (placar, sessao, contador)"]),
                (None,
                 ["Anote as colunas de cada tabela",
                  "Veja o ranking: top 5 por jogo"]),
                (["Abra http://127.0.0.1:8777/api/ranking"],
                 ["Compare JSON com saida do script"]),
                (["Proponha: Quem tem mais recordes?",
                  "Ou: Qual jogo foi mais jogado?"],
                 ["Abra Python e experimente queries"]),
            ]),
            ("ana",
             "SQLite e um caderno. Tables = paginas. Columns = colunas. "
             "Rows = linhas. Quando clica Registrar pontos, escreve uma linha. "
             "A query de ranking pede: 5 melhores de cada jogo — RANK() e "
             "PARTITION BY organizam."),
            ("cod",
             "servidor/arena.py (linhas 349-377) - schema:",
             SCHEMA),
            ("cod",
             "servidor/arena.py (linhas 546-578) - query de ranking:",
             RANKING),
            ("perg",
             ["Por que MAX(pontos) e nao pontos direto?",
              "O que PARTITION BY jogo_id faz?",
              "Por que WHERE posicao <= 5?"],
             ["Pega o maior recorde de cada jogador",
              "Separa rankings por jogo",
              "Mostra so os 5 primeiros"]),
            ("res",
             ["Identificou 3 tabelas (placar, sessao, contador)",
              "Entendeu MAX, GROUP BY, RANK, PARTITION BY",
              "JSON do navegador = query do script"]),
            ("saida",
             "Diferenca entre placar e sessao? Qual usa MAX e qual usa SUM?"),
        ]
    )

def aula4():
    gerar(
        "ATIVIDADE-ARENA-AULA4-FRONTEND1.docx",
        "ATIVIDADE: FRONT-END - HTML, CSS e JavaScript",
        "A estrutura da interface do ARENA",
        [
            ("obj",
             "Identificar estrutura HTML, seletores CSS e elementos "
             "dinamicos do JavaScript, relacionando-os aos componentes "
             "da interface do ARENA."),
            ("mat",
             ["Navegador (Chrome ou Edge)",
              "Acesso: http://127.0.0.1:8777",
              "Arquivo TESTAR_FRONTEND.bat"]),
            ("passos", [
                (["Abra TESTAR_FRONTEND.bat - mostra estatisticas"],
                 ["Anote: app.js 5054 linhas, index.html 1222"]),
                (None,
                 ["Abra http://127.0.0.1:8777",
                  "F12 -> aba Elements"]),
                (None,
                 ["Use a lupa para encontrar cada ID",
                  "Preencha a tabela: tag, classes, funcao JS"]),
            ]),
            ("ana",
             "HTML = moldura. CSS = tinta. JS = cerebro. No ARENA, tudo e "
             "JS puro — sem React, sem jQuery."),
            ("tab",
             "Elementos-chave do HTML",
             (["ID", "Tag HTML", "Classes", "Funcao JS"],
              [["#versao", "", "", "carregarEstado()"],
               ["#btnJogarSelecionado", "", "", "jogar()"],
               ["#rankingPontos", "", "", "desenharRanking()"],
               ["#listaJogos", "", "", "desenharJogos()"],
               ["#modalPlacar", "", "", "abrirPlacar()"],
               ["#chaveRede", "", "", "exibirChave()"],
               ["#perfilMiniNome", "", "", "carregarEstado()"]])),
            ("cod",
             "index.html (linhas 94-100) - estrutura principal:",
             HTML_SNIPPET),
            ("cod",
             "app.js (linhas 9, 39-40) - estado e utilidade:",
             ESTADO_JS),
            ("res",
             ["Identificou os 7 elementos na tela",
              "Entendeu: HTML=estrutura, CSS=aparencia, JS=comportamento",
              "Reconheceu fonte de verdade (estado) e funcao atalho ($)"]),
            ("saida",
             "Diferenca entre querySelector e createElement?"),
        ]
    )

def aula5():
    gerar(
        "ATIVIDADE-ARENA-AULA5-FRONTEND2.docx",
        "ATIVIDADE: FRONT-END - fetch() e API",
        "Como o navegador pede dados ao servidor",
        [
            ("obj",
             "Demonstrar o ciclo requisicao-resposta HTTP, analisar JSON e "
             "compreender assincronicidade com async/await."),
            ("mat",
             ["Navegador em: http://127.0.0.1:8777",
              "F12 -> aba Network e Console"]),
            ("passos", [
                (["Explique: browser faz pedidos HTTP ao servidor do ARENA"],
                 ["F12 -> aba Network",
                  "Recarregue (Ctrl+R)",
                  "Veja /api/estado, /api/catalogo"]),
                (None,
                 ["Clique botao Recordes",
                  "Network mostra /api/ranking",
                  "Clique duas vezes -> Response = JSON"]),
                (None,
                 ["No Console, cole e execute cada comando",
                  "Observe o JSON retornado"]),
            ]),
            ("cmd", API_COMANDOS),
            ("ana",
             "fetch() e um carteiro assincrono. await fetch() pausa ate "
             "retornar. Enquanto fora, o browser continua respondendo."),
            ("tab",
             "Rotas API do ARENA",
             (["Rota", "Metodo", "Retorna", "Quando"],
              ["/api/estado", "GET", "config, versao", "abertura"],
              ["/api/ranking", "GET", "placar + tempo", "botao Recordes"],
              ["/api/colegas", "GET", "lista de colegas", "rede local"],
              ["/api/conquistas", "GET", "conquistas", "perfil"],
              ["/api/placar", "POST", "ok + pendentes", "registrar pontos"],
              ["/api/sincronizar", "POST", "resultado", "enviar recordes"])),
            ("cod",
             "app.js (linhas 97-118) - funcao pedir():",
             PEDIR_JS),
            ("res",
             ["4 chamadas fetch() retornam JSON no console",
              "Na Network, viu pedidos em tempo real",
              "Entendeu: fetch=faz; await=espera",
              "Reconheceu /api/ranking browser = SQL do banco"]),
            ("saida",
             "Diferenca entre fetch e await? Qual falha se servidor offline?"),
        ]
    )

def aula6():
    gerar(
        "ATIVIDADE-ARENA-AULA6-FRONTEND3.docx",
        "ATIVIDADE: FRONT-END - DOM e Eventos",
        "Modificando a tela do ARENA em tempo real",
        [
            ("obj",
             "Manipular o DOM, registrar ouvintes de eventos e propor "
             "mudancas visuais na interface do ARENA."),
            ("mat",
             ["Navegador em: http://127.0.0.1:8777",
              "F12 -> aba Console"]),
            ("passos", [
                (["Explique: DOM e a arvore viva do HTML"],
                 ["F12 -> aba Console",
                  "Linha de comando no final"]),
                (None,
                 ["Execute cada teste (copie, cole, Enter)",
                  "Observe a mudanca na tela",
                  "Para voltar: recarregue (Ctrl+R)"]),
                (["Proporcione ajuda individualizada"],
                 ["Preencha a tabela de registro"]),
            ]),
            ("cmd", DOM_COMANDOS),
            ("tab",
             "Ficha de registro - DOM",
             (["Comando", "O que fez?", "Resultado", "Funcionou? (V/F)"],
              [["$('#versao').textContent", "leu texto", "", ""],
               ["$('#detalhesTitulo').textContent=...", "mudou texto", "", ""],
               ["$('#detalhesCapa').hidden=true", "escondeu capa", "", ""],
               ["document.createElement('p')", "criou paragrafo", "", ""],
               ["addEventListener('click',...)", "registro clique", "", ""]])),
            ("ana",
             "DOM = arvore de Lego digital. querySelector = pegar galho. "
             "createElement = recortar novo pedaco. addEventListener = "
             "colocar sino no galho."),
            ("cod",
             "app.js (linhas 4277-4280) - botao Jogar abre emulador:",
             EVENTOS_JS),
            ("info",
             "Ideia para rapidos:\n"
             "$('#detalhesJogo').style.background = '#e3f2fd'\n"
             "Voltar: $('#detalhesJogo').style.background = ''"),
            ("res",
             ["Leu textos com .textContent",
              "Modificou textos e cores com .style",
              "Escondeu/mostrou com .hidden",
              "Criou elementos com createElement",
              "Registrou eventos com addEventListener",
              "Entendeu: querySelector le, createElement cria"]),
            ("saida",
             "Diferenca entre querySelector e createElement? "
             "Por que usar DocumentFragment antes de replaceChildren?"),
        ]
    )


if __name__ == "__main__":
    limpar()
    print("Gerando atividades didaticas do ARENA...")
    aula1()
    aula2()
    aula3()
    aula4()
    aula5()
    aula6()
    print("\nTodos os 6 documentos .docx criados.")
